import urllib.request
import urllib.error
import json
import sys
import io
import openpyxl

BASE_URL = "http://127.0.0.1:5000"

def run_tests():
    print("--- TESTING EXPORT REPORT ENDPOINT ---")
    
    # 1. Admin Login
    print("1. Logging in as Admin...")
    req = urllib.request.Request(f"{BASE_URL}/api/admin/login", method="POST")
    req.add_header('Content-Type', 'application/json')
    req.data = json.dumps({"email": "admin@gmail.com", "password": "Admin@1234"}).encode('utf-8')
    
    try:
        with urllib.request.urlopen(req) as resp:
            data = json.loads(resp.read().decode('utf-8'))
            token = data.get("access_token")
            print("   [SUCCESS] Admin Token obtained.")
    except Exception as e:
        print(f"   [FAIL] Login failed: {e}")
        sys.exit(1)

    # 2. Test unauthorized request (no token)
    print("2. Testing Export Endpoint without token (Security Check)...")
    req_unauth = urllib.request.Request(f"{BASE_URL}/api/admin/reports/export?range=7d", method="GET")
    try:
        with urllib.request.urlopen(req_unauth) as resp:
            print("   [FAIL] Endpoint allowed request without token!")
    except urllib.error.HTTPError as e:
        if e.code in (401, 403):
            print(f"   [PASS] Endpoint properly returned {e.code} for unauthenticated request.")
        else:
            print(f"   [WARNING] Endpoint returned HTTP {e.code}")

    # 3. Test Export with Date Range filters: 7d, 30d, ytd, all
    for range_param in ["7d", "30d", "ytd", "all"]:
        url = f"{BASE_URL}/api/admin/reports/export?range={range_param}"
        print(f"3. Requesting Export Report with range='{range_param}'...")
        req_export = urllib.request.Request(url, method="GET")
        req_export.add_header('Authorization', f'Bearer {token}')
        
        try:
            with urllib.request.urlopen(req_export) as resp:
                status_code = resp.status
                headers = dict(resp.headers)
                content_type = headers.get("Content-Type", "")
                content_disp = headers.get("Content-Disposition", "")
                excel_bytes = resp.read()
                
                print(f"   [SUCCESS] Status: {status_code}")
                print(f"   Content-Type: {content_type}")
                print(f"   Content-Disposition: {content_disp}")
                print(f"   Downloaded File Size: {len(excel_bytes)} bytes")
                
                assert status_code == 200, f"Expected 200, got {status_code}"
                assert "spreadsheetml" in content_type, f"Invalid Content-Type {content_type}"
                assert "attachment" in content_disp, f"Missing attachment in Content-Disposition"
                
                # Parse with openpyxl to verify sheets
                wb = openpyxl.load_workbook(filename=io.BytesIO(excel_bytes))
                sheet_names = wb.sheetnames
                print(f"   Sheets found: {sheet_names}")
                
                expected_sheets = ["Dashboard Summary", "Bookings", "Customers", "Providers", "Revenue"]
                for sheet in expected_sheets:
                    assert sheet in sheet_names, f"Missing sheet: {sheet}"
                    ws = wb[sheet]
                    print(f"     - Sheet '{sheet}': {ws.max_row} rows x {ws.max_column} cols")
                
                print(f"   [PASS] Excel workbook structure and sheets validated for range='{range_param}'.")
        except Exception as e:
            print(f"   [FAIL] Error requesting export for '{range_param}': {e}")
            sys.exit(1)

    print("\nALL EXPORT REPORT BACKEND TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_tests()
